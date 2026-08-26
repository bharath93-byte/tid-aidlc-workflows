#!/usr/bin/env python3
"""Turn manual-test-cases.json into a formatted .xlsx.

Usage:
  python3 json_to_xlsx.py <path-to-manual-test-cases.json> [out.xlsx]

Sheet layout:
  1 story  → one sheet named after the ticket
  2+ stories → Summary (KPI cards + bar chart) + All Cases + one sheet per ticket
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY = "1B3A4B"
TEAL = "0D7377"
GOLD = "C9A227"
WHITE = "FFFFFF"
LIGHT = "F4F7F7"
ALT = "E8F1F1"
HIGH = "C0392B"
MED = "D68910"
LOW = "1E8449"
PASS = "1E8449"
FAIL = "C0392B"
WIP = "2471A3"
BLOCK = "7D3C98"
SKIP = "7F8C8D"
TODO = "B7950B"

PALETTE = [
    "1B4F72",
    "117A65",
    "6C3483",
    "B9770E",
    "1A5276",
    "922B21",
    "0E6655",
    "1B3A4B",
    "6E2C00",
    "154360",
]

HEADERS = [
    "Ticket",
    "Story",
    "Test Case ID",
    "Summary",
    "Description",
    "Type",
    "Priority",
    "Status",
    "Preconditions",
    "Test Steps",
    "Expected Results",
    "Verification Surfaces",
    "EARS Refs",
    "Demo Gaps",
    "Labels",
    "Tester",
    "Notes / Actual Result",
]

COL_WIDTHS = {
    "A": 14,
    "B": 42,
    "C": 18,
    "D": 48,
    "E": 52,
    "F": 16,
    "G": 12,
    "H": 14,
    "I": 46,
    "J": 52,
    "K": 52,
    "L": 28,
    "M": 22,
    "N": 28,
    "O": 28,
    "P": 16,
    "Q": 32,
}

THIN = Border(
    left=Side(style="thin", color="D0D5D5"),
    right=Side(style="thin", color="D0D5D5"),
    top=Side(style="thin", color="D0D5D5"),
    bottom=Side(style="thin", color="D0D5D5"),
)
HEADER_FILL = PatternFill("solid", fgColor=NAVY)
HEADER_FONT = Font(name="Calibri", bold=True, color=WHITE, size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(wrap_text=True, vertical="center", horizontal="center")


def ticket_keys(doc: dict) -> list[str]:
    return [k for k in doc if k != "_meta" and isinstance(doc[k], dict)]


def ticket_color(keys: list[str], key: str) -> str:
    return PALETTE[keys.index(key) % len(PALETTE)]


def fmt_list(items) -> str:
    if not items:
        return ""
    return "\n".join(f"• {x}" for x in items)


def steps_cols(steps) -> tuple[str, str]:
    actions, expected = [], []
    for s in steps or []:
        n = s.get("step", "")
        actions.append(f"{n}. {s.get('action', '')}")
        expected.append(f"{n}. {s.get('expected_result', '')}")
    return "\n".join(actions), "\n".join(expected)


def rows_for_ticket(doc: dict, key: str) -> list[list]:
    block = doc[key]
    story = block.get("title", "")
    out = []
    for tc in block.get("test_cases") or []:
        actions, expected = steps_cols(tc.get("steps", []))
        out.append(
            [
                key,
                story,
                tc.get("id", ""),
                tc.get("summary", ""),
                tc.get("description", ""),
                tc.get("type", ""),
                tc.get("priority", ""),
                "Not Started",
                fmt_list(tc.get("preconditions", [])),
                actions,
                expected,
                ", ".join(tc.get("verification_surfaces", [])),
                ", ".join(tc.get("ears_refs", [])),
                ", ".join(tc.get("demo_gap_refs", [])) or "—",
                ", ".join(tc.get("labels", [])),
                "",
                "",
            ]
        )
    return out


def style_header(ws, ncols: int) -> None:
    ws.row_dimensions[1].height = 24
    for col in range(1, ncols + 1):
        cell = ws.cell(1, col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = THIN
    ws.freeze_panes = "A2"


def apply_widths(ws) -> None:
    for col, w in COL_WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.oddHeader.left.text = "Manual Test Cases"
    ws.oddFooter.right.text = "Page &P of &N"


def fill_row(ws, row: int, values: list, color: str) -> None:
    pri = values[6]
    bg = ALT if row % 2 == 0 else LIGHT
    for i, val in enumerate(values, 1):
        cell = ws.cell(row, i, val)
        cell.alignment = CENTER if i in (1, 6, 7, 8) else WRAP
        cell.border = THIN
        cell.font = Font(name="Calibri", size=10, color="1C2833")
        cell.fill = PatternFill("solid", fgColor=bg)
    ws.cell(row, 1).fill = PatternFill("solid", fgColor=color)
    ws.cell(row, 1).font = Font(name="Calibri", bold=True, color=WHITE, size=10)
    pcell = ws.cell(row, 7)
    pcell.font = Font(name="Calibri", bold=True, color=WHITE, size=10)
    pcell.fill = PatternFill(
        "solid",
        fgColor=HIGH if pri == "High" else MED if pri == "Medium" else LOW,
    )
    scell = ws.cell(row, 8)
    scell.font = Font(name="Calibri", bold=True, color=WHITE, size=10)
    scell.fill = PatternFill("solid", fgColor=TODO)
    ws.row_dimensions[row].height = 72


def add_status_validation(ws, last_row: int) -> None:
    if last_row < 2:
        return
    dv = DataValidation(
        type="list",
        formula1='"Not Started,In Progress,Pass,Fail,Blocked,Skipped"',
        allow_blank=False,
        showDropDown=False,
        showErrorMessage=True,
        errorTitle="Invalid status",
        error="Pick a status from the list.",
    )
    dv.add(f"H2:H{last_row}")
    ws.add_data_validation(dv)
    for text, color in (
        ("Not Started", TODO),
        ("In Progress", WIP),
        ("Pass", PASS),
        ("Fail", FAIL),
        ("Blocked", BLOCK),
        ("Skipped", SKIP),
    ):
        ws.conditional_formatting.add(
            f"H2:H{last_row}",
            CellIsRule(
                operator="equal",
                formula=[f'"{text}"'],
                fill=PatternFill("solid", fgColor=color),
                font=Font(name="Calibri", bold=True, color=WHITE, size=10),
            ),
        )


def write_cases_sheet(ws, doc: dict, keys: list[str], only: str | None = None) -> int:
    ws.append(HEADERS)
    row_i = 2
    use = [only] if only else keys
    for key in use:
        color = ticket_color(keys, key)
        for values in rows_for_ticket(doc, key):
            fill_row(ws, row_i, values, color)
            row_i += 1
    last = max(row_i - 1, 1)
    style_header(ws, len(HEADERS))
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{last}"
    apply_widths(ws)
    add_status_validation(ws, last)
    ws.sheet_view.zoomScale = 90
    return last


def write_summary(ws, doc: dict, keys: list[str]) -> None:
    meta = doc.get("_meta") or {}
    epic = meta.get("epic", "")
    generated = meta.get("generated_at", "")
    sources = ", ".join(meta.get("sources") or [])
    total = sum(len((doc[k].get("test_cases") or [])) for k in keys)

    ws.sheet_view.showGridLines = False
    ws.merge_cells("B2:F2")
    ws["B2"] = f"{epic or 'Manual Test Cases'} — Manual Test Cases"
    ws["B2"].font = Font(name="Calibri", bold=True, color=NAVY, size=18)
    ws.merge_cells("B3:F3")
    ws["B3"] = f"Epic: {epic}   •   Generated: {generated}   •   Sources: {sources}"
    ws["B3"].font = Font(name="Calibri", color="4A5A5A", size=11)

    kpis = [
        (2, "Total Cases", total, NAVY),
        (3, "Stories", len(keys), TEAL),
        (4, "Not Started", total, TODO),
        (5, "Passed", 0, PASS),
        (6, "Failed", 0, FAIL),
    ]
    ws.row_dimensions[5].height = 18
    ws.row_dimensions[6].height = 28
    for col, label, value, color in kpis:
        cell_l = ws.cell(5, col, label)
        cell_l.font = Font(name="Calibri", bold=True, color=WHITE, size=10)
        cell_l.fill = PatternFill("solid", fgColor=color)
        cell_l.alignment = CENTER
        cell_v = ws.cell(6, col, value)
        cell_v.font = Font(name="Calibri", bold=True, color=color, size=20)
        cell_v.fill = PatternFill("solid", fgColor=LIGHT)
        cell_v.alignment = CENTER
        cell_l.border = THIN
        cell_v.border = THIN

    ws["B8"] = "Count by story"
    ws["B8"].font = Font(name="Calibri", bold=True, color=NAVY, size=13)

    sum_headers = ["Ticket", "Story", "Feature", "Cases", "High", "Medium", "Low", "Status"]
    for i, h in enumerate(sum_headers, 2):
        c = ws.cell(10, i, h)
        c.fill = HEADER_FILL
        c.font = HEADER_FONT
        c.alignment = CENTER
        c.border = THIN

    for r, key in enumerate(keys, 11):
        block = doc[key]
        cases = block.get("test_cases") or []
        highs = sum(1 for t in cases if t.get("priority") == "High")
        meds = sum(1 for t in cases if t.get("priority") == "Medium")
        lows = sum(1 for t in cases if t.get("priority") == "Low")
        vals = [
            key,
            block.get("title", ""),
            block.get("feature", ""),
            len(cases),
            highs,
            meds,
            lows,
            "Not Started",
        ]
        for i, v in enumerate(vals, 2):
            cell = ws.cell(r, i, v)
            cell.border = THIN
            cell.alignment = WRAP if i == 3 else CENTER
            cell.font = Font(name="Calibri", size=10)
            cell.fill = PatternFill("solid", fgColor=ALT if r % 2 == 0 else LIGHT)
        color = ticket_color(keys, key)
        ws.cell(r, 2).fill = PatternFill("solid", fgColor=color)
        ws.cell(r, 2).font = Font(name="Calibri", bold=True, color=WHITE, size=10)
        ws.cell(r, 9).fill = PatternFill("solid", fgColor=TODO)
        ws.cell(r, 9).font = Font(name="Calibri", bold=True, color=WHITE, size=10)
        ws.row_dimensions[r].height = 36

    last_data = 10 + len(keys)
    leg_row = last_data + 2
    ws.cell(leg_row, 2, "How to use").font = Font(
        name="Calibri", bold=True, color=NAVY, size=13
    )
    notes = [
        "Open All Cases (or a ticket tab). Filter by Ticket / Priority / Type / Status.",
        "Status column has a dropdown: Not Started, In Progress, Pass, Fail, Blocked, Skipped.",
        "Fill Tester and Notes / Actual Result as you execute. Default status is Not Started.",
        "Each row is one Jira-ready manual case (summary + steps + expected + live surfaces).",
    ]
    for i, note in enumerate(notes):
        ws.cell(leg_row + 1 + i, 2, f"• {note}").font = Font(
            name="Calibri", size=10, color="2C3E50"
        )
        ws.merge_cells(
            start_row=leg_row + 1 + i,
            start_column=2,
            end_row=leg_row + 1 + i,
            end_column=8,
        )

    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 56
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 10
    ws.column_dimensions["F"].width = 12
    ws.column_dimensions["G"].width = 10
    ws.column_dimensions["H"].width = 10
    ws.column_dimensions["I"].width = 16
    ws.row_dimensions[2].height = 28
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1

    chart = BarChart()
    chart.type = "col"
    chart.title = "Cases per story"
    chart.y_axis.title = "Cases"
    data = Reference(ws, min_col=5, min_row=10, max_row=last_data)
    cats = Reference(ws, min_col=2, min_row=11, max_row=last_data)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.legend = None
    chart.style = 10
    chart.y_axis.majorGridlines = None
    chart.height = 7
    chart.width = 15
    ws.add_chart(chart, f"B{leg_row + 7}")


def build_workbook(doc: dict) -> Workbook:
    keys = ticket_keys(doc)
    if not keys:
        raise SystemExit("No ticket keys found in JSON (expected objects besides _meta).")

    wb = Workbook()
    meta = doc.get("_meta") or {}
    wb.properties.title = f"{meta.get('epic', 'Manual Test Cases')} — Manual Test Cases"
    wb.properties.creator = "aidlc-manual-test-cases"
    wb.properties.subject = meta.get("epic", "")

    if len(keys) == 1:
        key = keys[0]
        ws = wb.active
        ws.title = key[:31]
        ws.sheet_properties.tabColor = ticket_color(keys, key)
        write_cases_sheet(ws, doc, keys, only=key)
        return wb

    summary = wb.active
    summary.title = "Summary"
    summary.sheet_properties.tabColor = GOLD
    write_summary(summary, doc, keys)

    all_ws = wb.create_sheet("All Cases")
    all_ws.sheet_properties.tabColor = TEAL
    write_cases_sheet(all_ws, doc, keys)

    for key in keys:
        ws_t = wb.create_sheet(key[:31])
        ws_t.sheet_properties.tabColor = ticket_color(keys, key)
        write_cases_sheet(ws_t, doc, keys, only=key)

    return wb


def main(argv: list[str]) -> int:
    if len(argv) < 2:
        print("Usage: json_to_xlsx.py <manual-test-cases.json> [out.xlsx]", file=sys.stderr)
        return 2
    src = Path(argv[1]).resolve()
    out = Path(argv[2]).resolve() if len(argv) > 2 else src.with_suffix(".xlsx")
    doc = json.loads(src.read_text(encoding="utf-8"))
    wb = build_workbook(doc)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    keys = ticket_keys(doc)
    mode = "single-sheet" if len(keys) == 1 else f"summary+all+{len(keys)}-tabs"
    print(f"Wrote {out} ({mode}; tickets={keys})")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
